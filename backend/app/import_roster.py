"""Import users and their groups from a roster spreadsheet.

Expected layout (any number of blocks, on the first sheet):
    Group 1241EA
    NrCrt | Nume                      | Grupa
    1     | POPESCU D. Ana-Maria      | 1241EA

Users are matched by full name, so re-running updates instead of duplicating.
Emails are generated as <first given name>.<surname>@<domain>.

Usage:  python -m app.import_roster attendance.xlsx [--domain D] [--password P] [--dry-run]
"""
import argparse
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Group, Role, User
from app.security import hash_password


def ascii_slug(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9-]+", "", text.lower().replace(" ", "-")).strip("-")


def split_name(full: str) -> tuple[str, str]:
    """'IONESCU L.C. Ștefan-Andrei' -> ('IONESCU', 'Ștefan-Andrei').
    Surname = leading all-caps words; father's initials (ending in '.') are skipped."""
    tokens = full.split()
    surname = []
    while tokens and tokens[0].isupper() and len(tokens[0].rstrip(".")) > 1 and not tokens[0].endswith("."):
        surname.append(tokens.pop(0))
    given = [t for t in tokens if not t.endswith(".")]
    return " ".join(surname) or full, " ".join(given)


def make_email(full: str, domain: str, taken: set[str]) -> str:
    surname, given = split_name(full)
    first = re.split(r"[\s-]+", given)[0] if given else ""
    local = ".".join(p for p in (ascii_slug(first), ascii_slug(surname)) if p) or "student"
    email, n = f"{local}@{domain}", 2
    while email in taken:
        email, n = f"{local}{n}@{domain}", n + 1
    taken.add(email)
    return email


def read_roster(path: Path) -> list[tuple[str, str]]:
    """Return (full_name, group) from either name/group or group/name blocks."""
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row[:3]]
        if len(cells) < 3 or not cells[0].isdigit():
            continue

        def group(value: str) -> str | None:
            match = re.fullmatch(r"(?:grupa?\s+)?(\d{4}[A-Za-z]{0,2})", value, re.I)
            return match.group(1).upper() if match else None

        if code := group(cells[2]):
            name = cells[1]
        elif code := group(cells[1]):
            name = cells[2]
        else:
            continue
        if name:
            rows.append((name, code))
    return rows


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--domain", default="student.example.com")
    parser.add_argument("--password", default="test", help="initial password for newly created users")
    parser.add_argument("--dry-run", action="store_true", help="print what would happen, change nothing")
    args = parser.parse_args()

    roster = read_roster(args.xlsx)
    if not roster:
        sys.exit("No users found in the spreadsheet")

    pw_hash = hash_password(args.password)
    created = updated = 0
    with SessionLocal() as db:
        groups = {g.name: g for g in db.scalars(select(Group))}
        taken = {e.lower() for e in db.scalars(select(User.email))}
        for full_name, group_name in roster:
            group = groups.get(group_name)
            if group is None:
                group = groups[group_name] = Group(name=group_name, description="Imported from roster")
                db.add(group)
            user = db.scalar(select(User).where(User.full_name == full_name))
            if user is None:
                user = User(full_name=full_name, email=make_email(full_name, args.domain, taken),
                            password_hash=pw_hash, role=Role.user)
                db.add(user)
                created += 1
                action = "create"
            else:
                updated += 1
                action = "update"
            if group not in user.groups:
                user.groups.append(group)
            print(f"{action:6} {group_name:8} {user.email:45} {full_name}")

        summary = {name: len([1 for _, g in roster if g == name]) for name in sorted({g for _, g in roster})}
        if args.dry_run:
            db.rollback()
            print(f"\nDry run: would create {created}, update {updated}. Groups: {summary}")
        else:
            db.commit()
            print(f"\nCreated {created}, updated {updated}. Groups: {summary}. "
                  f"New users' password: {args.password}")


if __name__ == "__main__":
    main()
